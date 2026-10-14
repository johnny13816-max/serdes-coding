"""Build an editable project-config draft from a labelled Ad Hoc snapshot."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable, Mapping, Optional

import numpy as np
from openpyxl import load_workbook

from ..adhoc import detect_adhoc_format, get_adhoc_parser
from ..adhoc.schema import AdhocConfigSnapshot, AdhocSetting, ProbeStatus
from .draft import ConversionDraft, MappingStatus, ProjectFieldDraft


def adhoc_to_project_config(
    adhoc_path: str | Path,
    output_path: str | Path,
    *,
    target_version: str,
    base_project_path: str | Path,
    format_id: Optional[str] = None,
    overrides: Optional[Mapping[tuple[str, str], Any]] = None,
    write_audit: bool = True,
) -> ConversionDraft:
    """Convert, validate, and write one external workbook in a single call.

    Use :func:`build_project_config_draft` instead when a human or GUI must
    inspect the mapping before export.
    """
    draft = build_project_config_draft(
        adhoc_path,
        target_version=target_version,
        base_project_path=base_project_path,
        format_id=format_id,
    )
    for (sheet, parameter), value in (overrides or {}).items():
        draft.set_override(sheet, parameter, value)

    # Local import avoids a module cycle: project_writer depends on the draft.
    from .project_writer import write_conversion_audit, write_project_config

    destination = write_project_config(draft, output_path)
    if write_audit:
        write_conversion_audit(draft, destination)
    return draft


def build_project_config_draft(
    adhoc_path: str | Path,
    *,
    target_version: str,
    base_project_path: str | Path,
    format_id: Optional[str] = None,
    package_selection: Optional[int] = None,
    tx_profile: Optional[str] = None,
    rx_profile: Optional[str] = None,
) -> ConversionDraft:
    """Parse one known source format and overlay it on a fixed project schema."""
    if target_version not in {"93A", "178A"}:
        raise ValueError(f"Unsupported target_version: {target_version!r}.")
    parser_selection = "automatic" if format_id is None else "manual"
    if format_id is None:
        probe = detect_adhoc_format(adhoc_path)
        if probe.status is not ProbeStatus.EXACT or probe.descriptor is None:
            raise ValueError(
                f"Ad Hoc format must be an exact automatic match; got {probe.status.value}. "
                "Pass format_id explicitly only after manual review."
            )
        format_id = probe.descriptor.format_id
        parser = get_adhoc_parser(format_id)
    else:
        parser = get_adhoc_parser(format_id)
        probe = parser.probe(adhoc_path)
        if probe.status is ProbeStatus.UNKNOWN:
            raise ValueError(f"Workbook is not compatible with manually selected parser {format_id}.")
    snapshot = parser.parse(adhoc_path)
    draft = ConversionDraft(
        source=snapshot,
        target_version=target_version,
        base_project_path=Path(base_project_path).expanduser().resolve(),
        fields=_read_project_schema(base_project_path),
        parser_selection=parser_selection,
        detection_status=probe.status.value,
        package_selection=package_selection,
        tx_profile=tx_profile,
        rx_profile=rx_profile,
    )
    _map_contract(snapshot, draft)
    return draft


def _map_contract(source, draft):
    from .mapping_rules import MappingContext, area_for, rules_for, array, SOURCE_EXTERNAL_CONTROLS
    rules=rules_for(draft.target_version)
    present={f.identity for f in draft.fields}
    for sheet,parameter in rules.keys()-present:
        draft.source_issues.append(f'Project schema missing required row: {sheet}.{parameter}')
    context=MappingContext(source,draft.package_selection,draft.tx_profile,draft.rx_profile)
    for item in draft.fields:
        area=area_for(item,draft.target_version)
        if area=='mapping':
            rule=rules[item.identity]
            context.reads=[]
            context.conditions=[]
            item.value=None
            item.source_key=rule.sources
            item.method=rule.method
            item.note=rule.description
            try:
                item.value=rule.evaluate(context)
                item.converted_value=item.value
                item.status=MappingStatus.EMPTY if item.value is None else (MappingStatus.DERIVED if rule.method=='derived' else MappingStatus.CONVERTED if rule.method!='direct' else MappingStatus.EXACT)
            except (KeyError,ValueError,TypeError,IndexError,OverflowError) as error:
                item.status=MappingStatus.MISSING
                item.note=f'{rule.description}: {error}'
            if context.reads:
                item.source_cell='; '.join(f'{s.sheet}!{s.value_cell}' for s in context.reads)
                item.source_block='; '.join(dict.fromkeys(s.block for s in context.reads if s.block)) or None
                item.source_value=context.reads[0].raw_value if len(context.reads)==1 else '; '.join(f'{s.key}={s.raw_value}' for s in context.reads)
                item.source_unit='; '.join(str(s.unit or '') for s in context.reads)
            if context.conditions:
                item.condition_key='; '.join(dict.fromkeys(s.key for s in context.conditions))
                item.condition_cell='; '.join(f'{s.sheet}!{s.value_cell}' for s in context.conditions)
                item.condition_value=(context.conditions[0].raw_value if len(context.conditions)==1
                                      else '; '.join(f'{s.key}={s.raw_value}' for s in context.conditions))
        elif area in {'policy','baseline'}:
            item.status=MappingStatus.FALLBACK
            item.method='project default'
            item.note='Project-owned setting; not an Ad Hoc mapping result'
        elif area=='case':
            if item.parameter=='case_id':
                item.status=MappingStatus.MANUAL
                item.method='case input'
                item.note='Case identity entered by the user; it must equal the case-folder name'
            elif item.parameter in {'source_config','case_source'}:
                item.value=str(source.source_path)
                item.converted_value=item.value
                item.source_key='Ad Hoc workbook path'
                item.source_value=str(source.source_path)
                item.status=MappingStatus.EXACT
                item.method='source provenance'
                item.note='Original Ad Hoc workbook used for this conversion'
            elif item.parameter=='source_note':
                item.value=f'{source.descriptor.format_id}; SHA-256 {source.source_sha256}'
                item.converted_value=item.value
                item.source_key='format_id; SHA-256'
                item.source_value=item.value
                item.status=MappingStatus.EXACT
                item.method='source provenance'
                item.note='Parser format identity and immutable source-workbook hash'
        elif area=='unregistered':
            item.status=MappingStatus.UNSUPPORTED
            item.note='No approved field contract exists'
    # Check known extra behavior, rather than pretending every parsed key is
    # a supported model input. Operational MATLAB controls are not model inputs.
    for key in ('c(-4)','c(-5)','c(2)'):
        for setting in source.find_all(key):
            try:
                if np.any(array(setting.value)!=0):
                    draft.source_issues.append(f'{key}: nonzero TX tap is unsupported by the project tap layout')
            except (ValueError,TypeError):draft.source_issues.append(f'{key}: cannot determine unsupported tap activity')
    if draft.target_version=='93A':
        for key in ('Include PCB','MLSE'):
            for setting in source.find_all(key):
                try:
                    if np.any(array(setting.value)!=0):draft.source_issues.append(f'{key}: active source feature is not represented by the 93A project schema')
                except (ValueError,TypeError):draft.source_issues.append(f'{key}: invalid source value')
    known=set(SOURCE_EXTERNAL_CONTROLS)|{'PKG_NAME','z_p select','Port Order','R_d','N_tc','trunc','c(-4)','c(-5)','c(2)'}
    for version in ('93A','178A'):
        for rule in rules_for(version).values():
            known.update(key.strip() for key in rule.sources.split(','))
    for setting in source.settings:
        if setting.key not in known:
            draft.source_issues.append(f'{setting.key}: source parameter has no reviewed mapping or exclusion rule')


def _read_project_schema(path: str | Path) -> list[ProjectFieldDraft]:
    # Project workbooks may have a deliberately cleared dimension record;
    # normal mode lets openpyxl calculate max_row/max_column reliably.
    workbook = load_workbook(path, read_only=False, data_only=False)
    fields: list[ProjectFieldDraft] = []
    try:
        for sheet_name in ("fixed_config", "search_config", "channels", "run_config"):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            headers = {
                str(sheet.cell(3, col).value).strip(): col
                for col in range(1, sheet.max_column + 1)
                if sheet.cell(3, col).value is not None
            }
            if sheet_name == "fixed_config":
                for row in range(4, sheet.max_row + 1):
                    parameter = sheet.cell(row, headers["Parameter"]).value
                    if parameter is None:
                        continue
                    parameter_class = (
                        sheet.cell(row, headers["Parameter Class"]).value
                        if "Parameter Class" in headers else "intrinsic"
                    )
                    fields.append(_field(
                        sheet_name, row, headers["Value"],
                        sheet.cell(row, headers["Domain"]).value if "Domain" in headers else "fixed",
                        parameter,
                        parameter_class,
                        sheet.cell(row, headers["Value"]).value,
                        sheet.cell(row, headers["Unit"]).value,
                        sheet.cell(row, headers["Description"]).value,
                    ))
            elif sheet_name == "search_config":
                for row in range(4, sheet.max_row + 1):
                    parameter = sheet.cell(row, headers["Parameter"]).value
                    if parameter is None:
                        continue
                    fields.append(_field(
                        sheet_name, row, headers["Values"], "search", parameter, "policy",
                        sheet.cell(row, headers["Values"]).value,
                        sheet.cell(row, headers["Unit"]).value,
                        sheet.cell(row, headers["Description"]).value,
                    ))
            elif sheet_name == "run_config":
                for row in range(4, sheet.max_row + 1):
                    profile = sheet.cell(row, headers["Profile"]).value
                    parameter = sheet.cell(row, headers["Parameter"]).value
                    if profile is None or parameter is None:
                        continue
                    name = f"{profile}.{parameter}"
                    fields.append(_field(
                        sheet_name, row, headers["Value"], profile, name, "policy",
                        sheet.cell(row, headers["Value"]).value,
                        sheet.cell(row, headers["Unit"]).value,
                        sheet.cell(row, headers["Description"]).value,
                    ))
            else:
                for row in range(4, sheet.max_row + 1):
                    kind = sheet.cell(row, headers["Kind"]).value
                    index = sheet.cell(row, headers["Index"]).value
                    if kind is None:
                        continue
                    row_id = f"{kind}[{index}]"
                    for column_name in ("S4P Path", "Port Order", "Use"):
                        name = f"{row_id}.{column_name}"
                        fields.append(_field(
                            sheet_name, row, headers[column_name], kind, name, "runtime",
                            sheet.cell(row, headers[column_name]).value, None,
                            sheet.cell(row, headers.get("Description", headers[column_name])).value,
                            status=MappingStatus.MANUAL,
                        ))
    finally:
        workbook.close()
    return fields


def _field(
    sheet: str, row: int, value_column: int, group: Any, parameter: Any,
    parameter_class: Any, value: Any, unit: Any, description: Any,
    *, status: MappingStatus = MappingStatus.FALLBACK,
) -> ProjectFieldDraft:
    return ProjectFieldDraft(
        sheet=sheet, row=row, value_column=value_column,
        group=str(group), parameter=str(parameter), parameter_class=str(parameter_class),
        value=value, unit=unit, description=description, base_value=value, status=status,
        note="Retained from the reviewed base project workbook.",
    )




def _array(value: Any) -> np.ndarray:
    if isinstance(value, (int, float, np.integer, np.floating)):
        return np.asarray([float(value)])
    text = str(value).strip()
    if text.startswith("[") and text.endswith("]"):
        text = text[1:-1].strip()
    rows: list[np.ndarray] = []
    for row_text in text.split(";"):
        parts = row_text.replace(",", " ").split()
        if not parts:
            continue
        if len(parts) == 1 and ":" in parts[0]:
            lo, step, hi = (float(item) for item in parts[0].split(":"))
            count = int(round((hi - lo) / step)) + 1
            rows.append(lo + step * np.arange(count))
        else:
            rows.append(np.asarray([float(item) for item in parts]))
    if not rows:
        raise ValueError(f"Cannot parse COM array: {value!r}.")
    return rows[0] if len(rows) == 1 else np.vstack(rows)


def _scalar(value: Any) -> float:
    values = _array(value).ravel()
    if len(values) != 1:
        raise ValueError(f"Expected scalar, got {value!r}.")
    return float(values[0])


def _bool(value: Any) -> bool:
    return bool(int(float(value)))


def _format_sequence(values: np.ndarray) -> str:
    return ",".join(f"{float(value):.15g}" for value in values)


def _row_text(value: Any, row: int, scale: float = 1.0) -> str:
    matrix = _array(value)
    selected = matrix if matrix.ndim == 1 else matrix[row, :]
    return _format_sequence(np.asarray(selected) * scale)


def _column_text(value: Any, column: int) -> str:
    matrix = _array(value)
    selected = matrix if matrix.ndim == 1 else matrix[:, column]
    return _format_sequence(np.asarray(selected))


def _range_text(value: Any) -> str:
    return _format_sequence(_array(value).ravel())


__all__ = ["adhoc_to_project_config", "build_project_config_draft"]
