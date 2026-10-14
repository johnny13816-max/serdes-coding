from pathlib import Path

from openpyxl import Workbook

from serdes_coding.io.adhoc import ProbeStatus, detect_adhoc_format, get_adhoc_parser
from serdes_coding.io.adhoc.ieee8023dj_d3p2_cakr_178_v1 import FORMAT_ID


def _write_minimum_exact_workbook(path: Path) -> None:
    parser = get_adhoc_parser(FORMAT_ID)
    workbook = Workbook()
    settings = workbook.active
    settings.title = "COM_Settings"
    for name in parser.descriptor.required_sheets:
        if name != "COM_Settings":
            workbook.create_sheet(name)

    for row, key in enumerate(parser.descriptor.signature_keys, start=1):
        settings.cell(row, 1, key)
        settings.cell(row, 2, row)
    row = 40
    for block in parser.descriptor.signature_blocks:
        settings.cell(row, 1, ".START")
        settings.cell(row, 2, block)
        settings.cell(row + 1, 1, "package_Z_c")
        settings.cell(row + 1, 2, "[100 100]")
        settings.cell(row + 2, 1, ".END")
        row += 4
    workbook.save(path)


def test_detect_and_parse_exact_format(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    _write_minimum_exact_workbook(path)

    probe = detect_adhoc_format(path)
    assert probe.status is ProbeStatus.EXACT
    assert probe.descriptor is not None
    assert probe.descriptor.format_id == FORMAT_ID

    snapshot = get_adhoc_parser(FORMAT_ID).parse(path)
    assert snapshot.descriptor.format_id == FORMAT_ID
    assert len(snapshot.source_sha256) == 64
    assert snapshot.require_one("f_b").value == 1
    assert snapshot.require_one("package_Z_c", block="PKG_Null").raw_value == "[100 100]"


def test_unknown_workbook_is_not_guessed(tmp_path: Path) -> None:
    path = tmp_path / "unknown.xlsx"
    Workbook().save(path)

    probe = detect_adhoc_format(path)
    assert probe.status is ProbeStatus.UNKNOWN
    assert probe.descriptor is None
