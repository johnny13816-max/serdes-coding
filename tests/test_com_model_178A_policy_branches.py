from dataclasses import fields, replace
from pathlib import Path
from shutil import copy2

import numpy as np
import openpyxl
import pytest

from serdes_coding.io.com_excel_io import excel_to_config_178A
from serdes_coding.models.com_model_178A import (
    COM,
    COMPolicyConfig,
    COMRunConfig,
    _resolve_adc_input_pmf_config,
    _resolve_final_com_pmf_config,
)


def _with_methods(cfg, **changes):
    return replace(cfg, policy=replace(cfg.policy, **changes))


def test_178a_branch_registry_covers_policy_and_run_fields():
    root = Path(__file__).resolve().parents[1]
    registry = (
        root / "docs" / "models" / "178A" / "branch_behavior_registry.md"
    ).read_text(encoding="utf-8")
    for field in fields(COMPolicyConfig):
        assert field.name in registry
    for field in fields(COMRunConfig):
        assert field.name in registry


@pytest.fixture(scope="module")
def reference_cfg():
    root = Path(__file__).resolve().parents[1]
    return excel_to_config_178A(
        str(root / "cases" / "sicopr_akinwale_vendorx_22db" / "178A" / "config.xlsx")
    )


def test_178a_out_of_measured_changes_only_scalar_vtf(reference_cfg):
    assert reference_cfg.policy.out_of_measured == "zero_padding"
    measured_cfg = _with_methods(reference_cfg, out_of_measured="cosine_decay")
    zero_paths = COM(reference_cfg).build_all_paths()
    measured_paths = COM(measured_cfg).build_all_paths()
    assert len(zero_paths) == len(measured_paths) == 6

    freqs = reference_cfg.link.freqs
    for zero, measured in zip(zero_paths, measured_paths):
        below = freqs <= zero.S_all.freqs[-1]
        above = ~below
        assert np.any(above)
        assert zero.kind == measured.kind
        np.testing.assert_array_equal(zero.S_ch.sdd, measured.S_ch.sdd)
        np.testing.assert_array_equal(zero.S_all.sdd, measured.S_all.sdd)
        np.testing.assert_allclose(zero.H_21.tf[below], measured.H_21.tf[below], rtol=0, atol=0)
        np.testing.assert_array_equal(zero.H_21.tf[above], np.zeros(np.count_nonzero(above)))


def test_178a_run_dy_override_bypasses_relative_and_absolute_grid(reference_cfg):
    As = reference_cfg.imp.R_LM / (reference_cfg.L - 1)
    coarse = replace(
        reference_cfg.execution.search_sweep,
        pmf_grid_quality="coarse",
        pmf_dy_override=2e-5,
    )
    resolved = _resolve_adc_input_pmf_config(reference_cfg.pmf, As, coarse)
    assert resolved.dy == 2e-5
    assert resolved.keep_mass == coarse.pmf_keep_mass
    final = _resolve_final_com_pmf_config(reference_cfg.pmf, As, coarse)
    assert final.dy == 2e-5
    assert final.keep_mass == 1.0

    old = replace(coarse, pmf_dy_override=None)
    assert _resolve_adc_input_pmf_config(reference_cfg.pmf, As, old).dy == pytest.approx(
        reference_cfg.pmf.dy_rel_As * As
    )
    assert _resolve_final_com_pmf_config(reference_cfg.pmf, As, old).dy == pytest.approx(
        min(reference_cfg.pmf.dy_abs_max, reference_cfg.pmf.dy_rel_As * As)
    )


@pytest.mark.parametrize("invalid", [0.0, -1e-5, np.nan, np.inf])
def test_178a_run_dy_override_rejects_nonpositive_or_nonfinite(invalid):
    with pytest.raises(ValueError, match="pmf_dy_override"):
        COMRunConfig(pmf_dy_override=invalid)


def test_178a_out_of_measured_rejects_unknown_method(reference_cfg):
    with pytest.raises(ValueError, match="out_of_measured"):
        _with_methods(reference_cfg, out_of_measured="other")


def test_178a_ctle_pulse_realization_defaults_and_rejects_unknown_method(reference_cfg):
    fields = COMPolicyConfig.__dataclass_fields__
    assert fields["ctle_pulse_realization"].default == "continuous_tf"
    assert fields["rxffe_record_realization"].default == "linear_convolution"
    with pytest.raises(ValueError, match="ctle_pulse_realization"):
        _with_methods(reference_cfg, ctle_pulse_realization="other")
    with pytest.raises(ValueError, match="rxffe_record_realization"):
        _with_methods(reference_cfg, rxffe_record_realization="other")


def test_178a_noise_grid_contract_and_reference_profile(reference_cfg):
    assert reference_cfg.policy.noise_grid_method == "fixed_ui"
    assert reference_cfg.policy.noise_grid_num_ui == 4096
    model = COM(reference_cfg)
    assert model.noise_link_cfg.sampled_nfft == 4096
    assert model.noise_link_cfg is not reference_cfg.link
    assert COMPolicyConfig().resolve_noise_grid(reference_cfg.link) is reference_cfg.link
    with pytest.raises(ValueError, match="even noise_grid_num_ui"):
        COMPolicyConfig(noise_grid_method="fixed_ui", noise_grid_num_ui=4095)
    with pytest.raises(ValueError, match="num_ui=None"):
        COMPolicyConfig(noise_grid_method="analysis_grid", noise_grid_num_ui=4096)


def test_178a_renormalization_policy_changes_only_channel_and_downstream(reference_cfg):
    assert reference_cfg.policy.channel_renormalization_order == "s4p_first"
    spec_cfg = _with_methods(reference_cfg, channel_renormalization_order="sdd_first")
    actual = COM(reference_cfg).build_all_paths()
    spec = COM(spec_cfg).build_all_paths()
    assert len(actual) == len(spec) == 6
    assert max(np.max(np.abs(a.S_ch.sdd - b.S_ch.sdd)) for a, b in zip(actual, spec)) > 1e-6
    for a, b in zip(actual, spec):
        np.testing.assert_array_equal(a.S_tx.sdd, b.S_tx.sdd)
        np.testing.assert_array_equal(a.S_rx.sdd, b.S_rx.sdd)
        np.testing.assert_array_equal(a.H_ffe.tf, b.H_ffe.tf)


def test_178a_renormalization_policy_rejects_unknown_method(reference_cfg):
    with pytest.raises(ValueError, match="channel_renormalization_order"):
        _with_methods(reference_cfg, channel_renormalization_order="unknown")


@pytest.mark.parametrize("pmf_method", ["gaussian_approx", "pmf_exact"])
def test_178a_pre_dte_adc_clipping_policy_changes_only_quantization_input(
    reference_cfg, pmf_method
):
    assert reference_cfg.policy.pre_dte_adc_clipping_convention == "matlab_4_17_reference_only"
    annex_cfg = _with_methods(reference_cfg, pre_dte_adc_clipping_convention="annex_d3")
    paths = COM(reference_cfg).build_all_paths()
    victim, xtalks = paths[0], paths[1:]
    model = COM(reference_cfg)
    common, h_xt_dsamp = model.calculate_pre_dte_imp_common(
        victim, [path.pulse.ir for path in xtalks]
    )
    run_cfg = replace(
        reference_cfg.execution.single_run,
        pre_dte_pmf_method=pmf_method,
        pmf_dy_override=1e-5,
        pmf_keep_mass=1.0,
    )
    reference = model.calculate_pre_dte_imp_at_pos(
        victim, 23, common, h_xt_dsamp, run_cfg
    )
    annex = COM(annex_cfg).calculate_pre_dte_imp_at_pos(
        victim, 23, common, h_xt_dsamp, run_cfg
    )

    ref_psd = reference.pre_dte.psd
    annex_psd = annex.pre_dte.psd
    assert reference.pre_dte.adc_input.clipping_convention == "matlab_4_17_reference_only"
    assert annex.pre_dte.adc_input.clipping_convention == "annex_d3"
    assert reference.pre_dte.adc_input.V_qc < annex.pre_dte.adc_input.V_qc
    if pmf_method == "pmf_exact":
        assert reference.pre_dte.adc_input.p_n is reference.pre_dte.adc_input.p_ga
        assert annex.pre_dte.adc_input.p_n is annex.pre_dte.adc_input.p_ga
    np.testing.assert_array_equal(reference.eq_ch.h_dsamp.ir, annex.eq_ch.h_dsamp.ir)
    for name in ("S_rn", "S_xn", "S_tn", "S_jn"):
        np.testing.assert_array_equal(getattr(ref_psd, name).psd, getattr(annex_psd, name).psd)
    np.testing.assert_allclose(
        ref_psd.S_total.psd - annex_psd.S_total.psd,
        ref_psd.S_qn.psd - annex_psd.S_qn.psd,
        rtol=0,
        atol=1e-25,
    )


def test_178a_pre_dte_adc_clipping_policy_rejects_unknown_method(reference_cfg):
    with pytest.raises(ValueError, match="pre_dte_adc_clipping_convention"):
        _with_methods(reference_cfg, pre_dte_adc_clipping_convention="unknown")


def test_178a_final_adc_quantization_policy_rejects_unknown_method(reference_cfg):
    assert reference_cfg.policy.final_adc_quantization_convention == "matlab_4_17_reference_only"
    with pytest.raises(ValueError, match="final_adc_quantization_convention"):
        _with_methods(reference_cfg, final_adc_quantization_convention="unknown")


def test_178a_align_to_ref_is_expanded_by_excel_adapter(reference_cfg, tmp_path):
    source = (
        Path(__file__).resolve().parents[1]
        / "cases" / "sicopr_akinwale_vendorx_22db" / "178A" / "config.xlsx"
    )
    workbook_path = tmp_path / "config_align_to_ref.xlsx"
    copy2(source, workbook_path)
    book = openpyxl.load_workbook(workbook_path)
    fixed = book["fixed_config"]
    raw_values = {
        "out_of_measured": "cosine_decay",
        "channel_renormalization_order": "sdd_first",
        "pre_dte_adc_clipping_convention": "annex_d3",
        "final_adc_quantization_convention": "annex_d3",
    }
    for row in fixed.iter_rows(min_row=4):
        name = row[1].value
        if name in raw_values:
            row[3].value = raw_values[name]
    fixed.append(["policy", "align_to_ref", "policy", True, "boolean", "reference preset"])
    book.save(workbook_path)

    aligned = excel_to_config_178A(str(workbook_path))
    policy = aligned.policy
    assert aligned.policy.applied_config_profile == "sicopr_reference"
    assert policy.out_of_measured == "zero_padding"
    assert policy.channel_renormalization_order == "s4p_first"
    assert policy.pre_dte_adc_clipping_convention == "matlab_4_17_reference_only"
    assert policy.final_adc_quantization_convention == "matlab_4_17_reference_only"
    assert policy.ctle_pulse_realization == "bilinear_iir"
    assert policy.rxffe_record_realization == "finite_record_circshift"
    assert policy.noise_grid_method == "fixed_ui"
    assert policy.noise_grid_num_ui == 4096
    assert policy.discrete_pmf_placement == "nearest"
    assert policy.interference_pmf_method == "pam_atoms"
    assert policy.gaussian_pmf_method == "reference_sample"
    assert policy.uniform_pmf_method == "nearest_centers"
    assert policy.pmf_scaling_method == "interpolate_ordinates"
    assert policy.added_noise_pdf == "scaled_total_pdf"
    assert policy.sequence_termination == "matlab_relative_tail"
    assert policy.reported_der == "baseline_pdf_re_evaluation"
    assert policy.gaussian_n_sigma_override == pytest.approx(3.540083799206174)
    assert policy.final_pmf_grid == "relative_As"
    assert policy.final_adc_pmf_projection == "reference_support"
    effective = COM(aligned)._effective_run_cfg(aligned.execution.search_sweep)
    assert effective.pre_dte_pmf_method == "pmf_exact"
    assert effective.pmf_grid_quality == "fine"
    assert effective.pmf_keep_mass == 1.0
    assert effective.pmf_dy_override is None
    final_grid = _resolve_final_com_pmf_config(
        aligned.pmf, aligned.imp.R_LM / (aligned.L - 1), effective,
        final_pmf_grid=policy.final_pmf_grid,
        gaussian_n_sigma_override=policy.gaussian_n_sigma_override,
    )
    assert final_grid.dy == pytest.approx(aligned.pmf.dy_rel_As * aligned.imp.R_LM / (aligned.L - 1))
    assert final_grid.gaussian_n_sigma == pytest.approx(policy.gaussian_n_sigma_override)
    exported = aligned.to_export_dict()["derived"]
    assert exported["policy"]["discrete_pmf_placement"] == "nearest"
    assert exported["policy"]["applied_config_profile"] == "sicopr_reference"
    overrides = exported["policy"]["config_profile_overrides"]
    assert overrides["policy.out_of_measured"] == {
        "raw": "cosine_decay",
        "effective": "zero_padding",
        "overridden_by_profile": "sicopr_reference",
    }
    assert overrides["policy.ctle_pulse_realization"] == {
        "raw": "continuous_tf",
        "effective": "bilinear_iir",
        "overridden_by_profile": "sicopr_reference",
    }
    assert overrides["policy.rxffe_record_realization"] == {
        "raw": "linear_convolution",
        "effective": "finite_record_circshift",
        "overridden_by_profile": "sicopr_reference",
    }
    assert overrides["policy.gaussian_pmf_method"] == {
        "raw": "bin_integral",
        "effective": "reference_sample",
        "overridden_by_profile": "sicopr_reference",
    }
    assert overrides["execution.search_sweep.pmf_grid_quality"] == {
        "raw": "coarse",
        "effective": "fine",
        "overridden_by_profile": "sicopr_reference",
    }

    direct = replace(
        reference_cfg,
        policy=replace(
            policy,
            applied_config_profile=None,
            config_profile_overrides={},
        ),
        execution=aligned.execution,
    )
    assert direct.policy == replace(
        aligned.policy,
        applied_config_profile=None,
        config_profile_overrides={},
    )
    assert direct.execution == aligned.execution


def test_178a_new_template_exposes_alignment_switch_and_noise_record_value():
    root = Path(__file__).resolve().parents[1]
    template = root / "cases" / "_templates" / "178A" / "config.xlsx"
    book = openpyxl.load_workbook(template, read_only=True, data_only=True)
    fields = {row[1].value for row in book["fixed_config"].rows if len(row) > 1}
    assert "align_to_ref" in fields
    value = next(row[3].value for row in book["fixed_config"].rows if len(row) > 3 and row[1].value == "align_to_ref")
    assert value is False
    assert "noise_grid_num_ui" in fields
    noise_value = next(
        row[3].value
        for row in book["fixed_config"].rows
        if len(row) > 3 and row[1].value == "noise_grid_num_ui"
    )
    assert noise_value is None
    for old_field in (
        "out_of_measured", "channel_renormalization_order",
        "pre_dte_adc_clipping_convention", "final_adc_quantization_convention",
    ):
        assert old_field not in fields
