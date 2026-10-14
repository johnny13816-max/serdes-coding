"""93A package coefficients must survive source mapping and reach both TLs."""
from dataclasses import asdict, replace
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from serdes_coding.models.com_model_93A import (
    COMPkgConfig, IEEECOMSparam, _build_txpkg_93A, _build_rxpkg_93A,
)
from serdes_coding.io.com_excel_io_93A import _fixed_pkg_config
from serdes_coding.io.conversion.adhoc_to_project import build_project_config_draft

ROOT = Path(__file__).resolve().parents[1]
NAMES = ("gamma_0", "a1", "a2", "tau")
PATHS = ("txpkg_victim", "txpkg_fext", "txpkg_next", "rxpkg")


@pytest.mark.parametrize("builder", [_build_txpkg_93A, _build_rxpkg_93A])
@pytest.mark.parametrize("lengths", [(7., None), (0., 3.), (7., 3.)])
def test_default_response_and_custom_coefficients(builder, lengths):
    f = np.linspace(0, 60e9, 101)
    cfg = COMPkgConfig(z_p=lengths[0], z_p2=lengths[1], Z_c=100., Z_c2=100.)
    total_length = sum(x for x in lengths if x is not None)
    old = IEEECOMSparam.pkg_trans_line_93A(f, 50., total_length, Zc=100.)
    np.testing.assert_allclose(builder(f, cfg).sdd, old.sdd, atol=1e-14)
    for name, value in zip(NAMES, (0.02, 0.003, 0.0003, 0.008)):
        custom = replace(cfg, **{name: value})
        # Matched TLs compose into exp(-gamma * total length), independently
        # of the package builder's implementation.
        ghz = f / 1e9
        gamma = np.full(f.shape, complex(custom.gamma_0))
        pos = ghz > 0
        gamma[pos] += custom.a1 * (1+1j) * np.sqrt(ghz[pos])
        gamma[pos] += (custom.a2 * (1-2j/np.pi*np.log(ghz[pos]))
                       + 2j*np.pi*custom.tau) * ghz[pos]
        actual = builder(f, custom).sdd
        np.testing.assert_allclose(actual[:, 1, 0], np.exp(-gamma*total_length), atol=1e-14)
        assert not np.allclose(actual, old.sdd)


@pytest.mark.parametrize("prefix", PATHS)
def test_adapter_requires_all_coefficients(prefix):
    cfg = COMPkgConfig(gamma_0=0.012, a1=0.003, a2=0.0004, tau=0.008)
    fixed = {f"{prefix}.{key}": value for key, value in asdict(cfg).items()}
    assert _fixed_pkg_config(fixed, prefix) == cfg
    for name in NAMES:
        key = f"{prefix}.{name}"
        for value in (None, ""):
            with pytest.raises((ValueError, TypeError, KeyError)):
                _fixed_pkg_config({**fixed, key: value}, prefix)
        missing = dict(fixed)
        del missing[key]
        with pytest.raises(KeyError, match=name):
            _fixed_pkg_config(missing, prefix)


def test_template_and_source_mapping():
    template = ROOT / "cases/_templates/93A/config.xlsx"
    with_template = load_workbook(template, read_only=True)
    rows = {r[1]: r for r in with_template['fixed_config'].iter_rows(min_row=4, values_only=True) if r[1]}
    for prefix in PATHS:
        for name in NAMES:
            row = rows[f"{prefix}.{name}"]
            assert row[2] == 'intrinsic' and row[3] is None
    with_template.close()
    source = ROOT / 'reference_data/COM_adhoc_config/sicopr/config_com_dj_200G_CAKR_178_PKGA_06_2_2025__Case1.xlsx'
    if not source.exists():
        pytest.skip('Local external reference workbook is not available')
    draft = build_project_config_draft(source, target_version='93A', base_project_path=template)
    for prefix in PATHS:
        fields = [draft.get('fixed_config', f'{prefix}.{name}') for name in NAMES]
        gamma = draft.source.require_one('package_tl_gamma0_a1_a2', block=fields[0].source_block)
        from serdes_coding.io.conversion.adhoc_to_project import _array
        np.testing.assert_allclose([f.value for f in fields[:3]], _array(gamma.value).ravel())
        assert fields[3].value == float(draft.source.require_one('package_tl_tau', block=fields[3].source_block).value)
