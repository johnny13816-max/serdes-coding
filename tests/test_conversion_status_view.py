import unittest
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from serdes_coding.interfaces.conversion_status import status_group
from serdes_coding.io.conversion.draft import ConversionDraft, MappingStatus, ProjectFieldDraft
from serdes_coding.io.conversion.paths import case_config_path, validate_case_id
from serdes_coding.io.conversion.adhoc_to_project import _read_project_schema
from serdes_coding.io.conversion.adhoc_to_project import build_project_config_draft
from serdes_coding.io.conversion.existing_project import (
    load_existing_project,
    save_existing_project,
    validate_existing_project,
)
from serdes_coding.interfaces.com_config_conversion_gui import _validated_source_path, COM_ADHOC_ROOT


class TestConversionStatusView(unittest.TestCase):
    def test_detailed_statuses_group_into_distinct_user_states(self):
        expected = {
            MappingStatus.EXACT: "success",
            MappingStatus.CONVERTED: "success",
            MappingStatus.DERIVED: "success",
            MappingStatus.FALLBACK: "success",
            MappingStatus.EMPTY: "empty",
            MappingStatus.MANUAL: "fail",
            MappingStatus.MISSING: "fail",
            MappingStatus.UNSUPPORTED: "fail",
            MappingStatus.OVERRIDE: "success",
        }
        self.assertEqual({status_group(k) for k in expected}, set(expected.values()))
        for detailed, display in expected.items():
            self.assertEqual(status_group(detailed), display)

    def test_override_keeps_the_previous_mapping_status(self):
        field = ProjectFieldDraft(
            sheet="fixed_config", row=4, value_column=3, group="link",
            parameter="fb", parameter_class="intrinsic", value=1.0, unit="Hz",
            description="baud rate", base_value=1.0,
            status=MappingStatus.CONVERTED,
        )
        draft = ConversionDraft(
            source=None, target_version="178A", base_project_path=None,
            fields=[field], parser_selection="automatic", detection_status="exact",
        )
        draft.set_override("fixed_config", "fb", 2.0)
        self.assertIs(field.status, MappingStatus.OVERRIDE)
        self.assertIs(field.original_status, MappingStatus.CONVERTED)

    def test_case_id_is_a_safe_folder_component_for_both_versions(self):
        root = Path("C:/test-project")
        for version in ("93A", "178A"):
            self.assertEqual(
                case_config_path(root, "case_260915", version),
                root / "cases" / "case_260915" / version / "config.xlsx",
            )
        for invalid in ("", "..", "../outside", "case/name", "C:\\outside"):
            with self.assertRaises(ValueError):
                validate_case_id(invalid)

    def test_templates_do_not_expose_internal_pmf_thresholds_as_project_fields(self):
        project_root = Path(__file__).resolve().parents[1]
        forbidden = {"dy_override", "tap_abs_th_override"}
        for version in ("93A", "178A"):
            fields = _read_project_schema(
                project_root / "cases" / "_templates" / version / "config.xlsx"
            )
            self.assertFalse(forbidden.intersection(field.parameter for field in fields))

    def test_schema_reader_accepts_both_project_workbook_layouts(self):
        project_root = Path(__file__).resolve().parents[1]
        for version in ("93A", "178A"):
            fields = _read_project_schema(project_root / "cases" / "_templates" / version / "config.xlsx")
            self.assertIn(("fixed_config", "case_id"), {field.identity for field in fields})
            self.assertTrue(any(field.sheet == "channels" for field in fields))

    def test_178a_template_contains_shared_required_pmf_defaults(self):
        project_root = Path(__file__).resolve().parents[1]
        workbook = load_workbook(project_root / "cases" / "_templates" / "178A" / "config.xlsx", data_only=True)
        try:
            sheet = workbook["fixed_config"]
            headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
            values = {
                sheet.cell(row, headers["Parameter"]).value: sheet.cell(row, headers["Value"]).value
                for row in range(4, sheet.max_row + 1)
            }
            self.assertEqual(values["dy_rel_As"], 0.001)
            self.assertEqual(values["dy_abs_max"], 1e-5)
            self.assertEqual(values["tap_rel_As"], 0.001)
            self.assertEqual(values["gaussian_n_sigma"], 8)
        finally:
            workbook.close()

    def test_178a_template_tx_main_cursor_matches_project_tap_vector(self):
        project_root = Path(__file__).resolve().parents[1]
        workbook = load_workbook(project_root / "cases" / "_templates" / "178A" / "config.xlsx", data_only=True)
        try:
            sheet = workbook["fixed_config"]
            headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
            row = next(
                row for row in range(4, sheet.max_row + 1)
                if sheet.cell(row, headers["Parameter"]).value == "num_pre"
            )
            self.assertEqual(sheet.cell(row, headers["Value"]).value, 3)
            self.assertEqual(sheet.cell(row, headers["Parameter Class"]).value, "derived")
        finally:
            workbook.close()

    def test_178a_mapping_does_not_use_rx_pre_tap_length_as_tx_main_index(self):
        project_root = Path(__file__).resolve().parents[1]
        source = (
            project_root / "reference_data" / "COM_adhoc_config" / "hansel"
            / "config_com_dj_200G_CAKR_178_PKGA_06_2_2025__Case1_TXLEOn.xlsx"
        )
        draft = build_project_config_draft(
            source,
            target_version="178A",
            base_project_path=project_root / "cases" / "_templates" / "178A" / "config.xlsx",
        )
        self.assertEqual(draft.source.require_one("ffe_pre_tap_len").value, 6)
        self.assertEqual(draft.get("fixed_config", "num_pre").value, 3)
        self.assertIs(draft.get("fixed_config", "num_pre").status, MappingStatus.DERIVED)

    def test_178a_missing_mlsd_trunc_is_an_error_but_delta_com_an_may_be_empty(self):
        project_root = Path(__file__).resolve().parents[1]
        source = (
            project_root / "reference_data" / "COM_adhoc_config" / "hansel"
            / "config_com_dj_200G_CAKR_178_PKGA_06_2_2025__Case1_TXLEOn.xlsx"
        )
        draft = build_project_config_draft(
            source,
            target_version="178A",
            base_project_path=project_root / "cases" / "_templates" / "178A" / "config.xlsx",
        )
        trunc = draft.get("fixed_config", "mlsd_trunc_len")
        delta = draft.get("fixed_config", "delta_com_an")
        self.assertIsNone(trunc.value)
        self.assertIs(trunc.status, MappingStatus.MISSING)
        self.assertIsNone(delta.value)
        self.assertIs(delta.status, MappingStatus.EMPTY)
        errors = draft.validation_errors()
        self.assertTrue(any("mlsd_trunc_len: missing" in error for error in errors))
        self.assertFalse(any("delta_com_an" in error for error in errors))

    def test_93a_fixed_filter_baseline_defaults_to_zero(self):
        project_root = Path(__file__).resolve().parents[1]
        workbook = load_workbook(project_root / "cases" / "_templates" / "93A" / "config.xlsx", data_only=True)
        try:
            sheet = workbook["fixed_config"]
            headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
            values = {
                sheet.cell(row, headers["Parameter"]).value: sheet.cell(row, headers["Value"]).value
                for row in range(4, sheet.max_row + 1)
            }
            for name in ("c_m2", "c_m1", "c_1", "g_DC", "g_DC2"):
                self.assertEqual(values[name], 0)
        finally:
            workbook.close()

    def test_gui_source_path_must_be_inside_project_reference_data(self):
        valid = next(COM_ADHOC_ROOT.rglob("*.xlsx"))
        self.assertEqual(_validated_source_path(str(valid)), valid.resolve())
        with self.assertRaises(ValueError):
            _validated_source_path(str(Path(__file__).resolve()))

    def test_open_existing_config_detects_both_project_versions(self):
        project_root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            for version in ("93A", "178A"):
                target = root / "cases" / f"open_{version}" / version / "config.xlsx"
                target.parent.mkdir(parents=True)
                shutil.copy2(project_root / "cases" / "_templates" / version / "config.xlsx", target)
                draft = load_existing_project(target, root)
                self.assertEqual(draft.target_version, version)
                self.assertIn(("fixed_config", "case_id"), {field.identity for field in draft.fields})

    def test_open_existing_config_keeps_fields_visible_when_folder_metadata_mismatches(self):
        project_root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            target = root / "cases" / "legacy_case" / "178A" / "config.xlsx"
            target.parent.mkdir(parents=True)
            shutil.copy2(project_root / "cases" / "_templates" / "93A" / "config.xlsx", target)
            workbook = load_workbook(target)
            sheet = workbook["fixed_config"]
            headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
            case_row = next(
                row for row in range(4, sheet.max_row + 1)
                if sheet.cell(row, headers["Parameter"]).value == "case_id"
            )
            sheet.cell(case_row, headers["Value"], "older_case_id")
            workbook.save(target)
            workbook.close()

            draft = load_existing_project(target, root)
            self.assertEqual(draft.target_version, "93A")
            self.assertGreater(len(draft.fields), 0)
            self.assertEqual(len(draft.open_warnings), 2)
            self.assertTrue(any("folder" in item.lower() for item in validate_existing_project(draft, target)))

    def test_existing_config_edits_preserve_runtime_readability_and_record_override(self):
        project_root = Path(__file__).resolve().parents[1]
        source_dir = project_root / "cases" / "sicopr_akinwale_vendorx_22db" / "178A"
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            target_dir = root / "cases" / "sicopr_edit_test" / "178A"
            target_dir.mkdir(parents=True)
            shutil.copy2(source_dir / "config.xlsx", target_dir / "config.xlsx")
            shutil.copy2(source_dir / "config.mapping.csv", target_dir / "config.mapping.csv")
            shutil.copy2(source_dir / "config.conversion.json", target_dir / "config.conversion.json")
            workbook = load_workbook(target_dir / "config.xlsx")
            sheet = workbook["fixed_config"]
            headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
            case_row = next(
                row for row in range(4, sheet.max_row + 1)
                if sheet.cell(row, headers["Parameter"]).value == "case_id"
            )
            sheet.cell(case_row, headers["Value"], "sicopr_edit_test")
            workbook.save(target_dir / "config.xlsx")
            workbook.close()

            draft = load_existing_project(target_dir / "config.xlsx", root)
            field = draft.get("fixed_config", "fb")
            expected_fb = field.value + 1e9
            draft.set_override("fixed_config", "fb", expected_fb)
            self.assertEqual(validate_existing_project(draft, target_dir / "config.xlsx"), ())
            saved = save_existing_project(draft, target_dir / "config.xlsx", root)

            reopened = load_existing_project(saved, root)
            self.assertEqual(reopened.get("fixed_config", "fb").value, expected_fb)
            self.assertIs(reopened.get("fixed_config", "fb").status, MappingStatus.OVERRIDE)
            self.assertEqual(reopened.get("fixed_config", "fb").original_status, MappingStatus.CONVERTED)

            new_case_id = "sicopr_save_as_test"
            new_case_root = root / "cases" / new_case_id
            new_case_root.mkdir()
            (new_case_root / "case_manifest.yaml").write_text(f"case_id: {new_case_id}\n", encoding="utf-8")
            reopened.set_override("fixed_config", "case_id", new_case_id)
            reopened.case_id = new_case_id
            new_path = new_case_root / "178A" / "config.xlsx"
            saved_as = save_existing_project(reopened, new_path, root)
            self.assertTrue(saved_as.is_file())
            self.assertTrue(saved_as.with_suffix(".conversion.json").is_file())
            self.assertEqual(load_existing_project(saved_as, root).case_id, new_case_id)


if __name__ == "__main__":
    unittest.main()
