"""Load, inspect, validate and save existing project configuration workbooks."""

from __future__ import annotations

import csv
import json
import os
import shutil
import tempfile
from dataclasses import dataclass, field, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .adhoc_to_project import _read_project_schema
from .draft import MappingStatus, ProjectFieldDraft
from .paths import case_config_path, validate_case_id


@dataclass
class ExistingProjectDraft:
    """Editable snapshot of a project workbook, independent of Ad Hoc input."""

    source_path: Path
    target_version: str
    case_id: str
    fields: list[ProjectFieldDraft]
    audit_rows: dict[tuple[str, str], dict[str, str]] = field(default_factory=dict)
    changed: set[tuple[str, str]] = field(default_factory=set)
    open_warnings: tuple[str, ...] = ()

    def get(self, sheet: str, parameter: str) -> ProjectFieldDraft:
        matches = [item for item in self.fields if item.identity == (sheet, parameter)]
        if len(matches) != 1:
            raise KeyError(f"Expected one project field {sheet}.{parameter}, got {len(matches)}.")
        return matches[0]

    def set_override(self, sheet: str, parameter: str, value: Any) -> None:
        item = self.get(sheet, parameter)
        if value == item.value:
            return
        if item.status is not MappingStatus.OVERRIDE:
            item.original_status = item.status
        item.value = value
        item.status = MappingStatus.OVERRIDE
        self.changed.add(item.identity)


def load_existing_project(path: str | Path, project_root: str | Path) -> ExistingProjectDraft:
    """Read a canonical case workbook and its optional conversion audit."""
    root = Path(project_root).expanduser().resolve(strict=True)
    source = Path(path).expanduser().resolve(strict=True)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("Select an existing .xlsx project config.")
    cases_root = (root / "cases").resolve(strict=True)
    try:
        relative = source.relative_to(cases_root)
    except ValueError as error:
        raise ValueError(f"Project configs must be inside {cases_root}.") from error
    if len(relative.parts) != 3 or relative.parts[-1].lower() != "config.xlsx":
        raise ValueError("Open a case config at cases/<case_id>/<93A|178A>/config.xlsx.")
    case_id, folder_version, _ = relative.parts
    validate_case_id(case_id)
    if folder_version not in {"93A", "178A"}:
        raise ValueError("Config folder must be named 93A or 178A.")

    open_warnings: list[str] = []
    workbook = load_workbook(source, read_only=False, data_only=True)
    try:
        required = {"fixed_config", "search_config", "channels"}
        if not required.issubset(workbook.sheetnames):
            raise ValueError("Workbook is missing fixed_config, search_config, or channels sheet.")
        version = "178A" if "run_config" in workbook.sheetnames else "93A"
        if version != folder_version:
            open_warnings.append(
                f"Workbook layout is {version}, but it is stored under the {folder_version} folder."
            )
        fixed = workbook["fixed_config"]
        headers = {
            str(fixed.cell(3, col).value).strip(): col
            for col in range(1, fixed.max_column + 1)
            if fixed.cell(3, col).value is not None
        }
        case_row = next(
            (row for row in range(4, fixed.max_row + 1)
             if fixed.cell(row, headers.get("Parameter", 0)).value == "case_id"),
            None,
        )
        if case_row is None or "Value" not in headers:
            raise ValueError("Workbook has no fixed_config.case_id field.")
        workbook_case_id = str(fixed.cell(case_row, headers["Value"]).value or "").strip()
    finally:
        workbook.close()
    if workbook_case_id and workbook_case_id != case_id:
        open_warnings.append(
            f"fixed_config.case_id is {workbook_case_id!r}, but folder name is {case_id!r}."
        )

    fields = _read_project_schema(source)
    audit_path = source.with_suffix(".mapping.csv")
    audit_rows: dict[tuple[str, str], dict[str, str]] = {}
    if audit_path.is_file():
        with audit_path.open("r", newline="", encoding="utf-8-sig") as stream:
            for row in csv.DictReader(stream):
                audit_rows[(row.get("sheet", ""), row.get("parameter", ""))] = row
        for item in fields:
            row = audit_rows.get(item.identity)
            if not row:
                continue
            try:
                item.status = MappingStatus(row.get("status", "manual"))
            except ValueError:
                item.status = MappingStatus.MANUAL
            try:
                item.original_status = MappingStatus(row.get("original_status", ""))
            except ValueError:
                item.original_status = None
            item.source_key = row.get("source_key") or None
            item.source_block = row.get("source_block") or None
            item.source_cell = row.get("source_cell") or None
            item.source_value = row.get("source_value") or None
            item.note = row.get("note") or ""
            item.source_unit = row.get('source_unit')
            item.converted_value = row.get('converted_value')
            item.method = row.get('method') or ''
            item.condition_key = row.get('condition_key') or None
            item.condition_cell = row.get('condition_cell') or None
            item.condition_value = row.get('condition_value') or None
            if item.status is MappingStatus.OVERRIDE:
                item.base_value = row.get("base_value") or item.base_value

    return ExistingProjectDraft(source, version, case_id, fields, audit_rows, open_warnings=open_warnings)


def canonical_destination(project_root: str | Path, case_id: str, version: str) -> Path:
    """Return the only supported save location for a project workbook."""
    return case_config_path(project_root, case_id, version)


def validate_existing_project(draft: ExistingProjectDraft, destination: str | Path) -> tuple[str, ...]:
    """Check case identity, enabled channel inputs and the versioned runtime parser."""
    target = Path(destination).expanduser().resolve()
    errors: list[str] = []
    from .mapping_rules import rules_for, outcome, area_for
    values={item.identity:item.value for item in draft.fields}
    for identity in rules_for(draft.target_version).keys()-values.keys():
        errors.append(f'Project schema missing required row: {identity[0]}.{identity[1]}')
    for item in draft.fields:
        if area_for(item,draft.target_version) in {'mapping','policy','baseline'}:
            # Validate the adopted project value, independently of old source
            # audit status. An existing file cannot prove source mapping anew.
            if outcome(replace(item,status=MappingStatus.OVERRIDE),draft.target_version,values)=='fail':
                errors.append(f'{item.sheet}.{item.parameter}: invalid or missing project value')
    try:validate_case_id(draft.case_id)
    except ValueError as error:errors.append(str(error))
    if target.parent.parent.name != draft.case_id:
        errors.append(
            f"Case ID {draft.case_id!r} must match destination folder {target.parent.parent.name!r}."
        )
    if target.parent.name != draft.target_version or target.name.lower() != "config.xlsx":
        errors.append("Destination must be cases/<case_id>/<version>/config.xlsx.")
    try:
        workbook_case_id = str(draft.get("fixed_config", "case_id").value or "").strip()
        if workbook_case_id != draft.case_id:
            errors.append(
                f"fixed_config.case_id {workbook_case_id!r} must match case folder {draft.case_id!r}."
            )
    except KeyError:
        errors.append("Workbook is missing fixed_config.case_id.")
    validation_fields=[]
    for item in draft.fields:
        validation_fields.append(item)
        if item.sheet != "channels" or not item.parameter.endswith(".S4P Path"):
            continue
        row_id = item.parameter.rsplit(".", 1)[0]
        enabled = next((field.value for field in draft.fields
                        if field.sheet == "channels" and field.parameter == f"{row_id}.Use"), True)
        if isinstance(enabled, str):
            enabled = enabled.strip().lower() not in {"false", "0", "no", "off", ""}
        if not bool(enabled):
            continue
        value = str(item.value or "").strip()
        if not value:
            errors.append(f"{item.parameter} is required for an enabled channel row.")
            continue
        from ..com_excel_common import _resolve_channel_path
        channel_path = Path(_resolve_channel_path(target, value)).expanduser()
        validation_fields[-1]=replace(item,value=str(channel_path.resolve()))
        if not channel_path.is_file():
            errors.append(f"{item.parameter} does not resolve to a file: {channel_path}")
    if errors:
        return tuple(errors)

    # Validate workbook-to-runtime parsing without constructing or running a COM model.
    fd, temp_name = tempfile.mkstemp(prefix=".config_check_", suffix=".xlsx")
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        _write_values(draft.source_path, temp_path, validation_fields)
        try:
            if draft.target_version == "178A":
                from ..com_excel_io_178A import excel_to_config_178A, excel_to_search_config_178A
                excel_to_config_178A(str(temp_path))
                excel_to_search_config_178A(str(temp_path))
            else:
                from ..com_excel_io_93A import excel_to_config_93A, excel_to_search_config_93A
                excel_to_config_93A(str(temp_path))
                excel_to_search_config_93A(str(temp_path))
        except Exception as error:
            errors.append(f"Runtime config validation failed: {type(error).__name__}: {error}")
    finally:
        temp_path.unlink(missing_ok=True)
    return tuple(errors)


def save_existing_project(
    draft: ExistingProjectDraft,
    destination: str | Path,
    project_root: str | Path,
) -> Path:
    """Save edited cells while preserving workbook structure and write an edit audit."""
    target = canonical_destination(project_root, draft.case_id, draft.target_version)
    requested = Path(destination).expanduser().resolve()
    if requested != target:
        raise ValueError(f"Project configs must be saved to the canonical path: {target}")
    errors = validate_existing_project(draft, target)
    if errors:
        raise ValueError("Cannot save invalid project config:\n- " + "\n- ".join(errors))
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=".config_edit_", suffix=".xlsx", dir=target.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        _write_values(draft.source_path, temp_path, draft.fields)
        os.replace(temp_path, target)
    finally:
        temp_path.unlink(missing_ok=True)
    if target != draft.source_path:
        source_conversion_audit = draft.source_path.with_suffix(".conversion.json")
        if source_conversion_audit.is_file():
            shutil.copy2(source_conversion_audit, target.with_suffix(".conversion.json"))
    _write_edit_audit(draft, target)
    draft.source_path = target
    draft.changed.clear()
    return target


def _write_values(source: Path, destination: Path, fields: list[ProjectFieldDraft]) -> None:
    workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        from .channels import write_fields
        write_fields(workbook, fields)
        workbook.save(destination)
    finally:
        workbook.close()


def _write_edit_audit(draft: ExistingProjectDraft, target: Path) -> None:
    if not draft.changed:
        return
    audit_path = target.with_suffix(".mapping.csv")
    columns = (
        "sheet", "group", "parameter", "parameter_class", "status", "original_status", "value", "unit",
        "source_key", "source_block", "source_cell", "source_value", "base_value", "note",
        "source_unit", "converted_value", "method", "condition_key", "condition_cell",
        "condition_value",
    )
    identities={item.identity for item in draft.fields}
    rows = {key:row for key,row in draft.audit_rows.items() if key in identities}
    for item in draft.fields:
        if item.identity not in draft.changed and not (item.sheet=='channels' and ('channels','__rows__') in draft.changed):
            continue
        previous = rows.get(item.identity, {})
        rows[item.identity] = {
            **previous,
            "sheet": item.sheet,
            "group": item.group,
            "parameter": item.parameter,
            "parameter_class": item.parameter_class,
            "status": "override",
            "original_status": (
                previous.get("original_status") or "existing_untracked"
                if previous.get("status") == "override"
                else previous.get("status", "existing_untracked")
            ),
            "value": str(item.value),
            "unit": str(item.unit or ""),
            "source_key": item.source_key or "",
            "source_block": item.source_block or "",
            "source_cell": item.source_cell or "",
            "source_value": str(item.source_value) if item.source_value is not None else "",
            "base_value": previous.get("base_value", str(item.base_value)),
            "note": item.note,
            "source_unit": item.source_unit,
            "converted_value": item.converted_value,
            "method": item.method,
            "condition_key": item.condition_key or "",
            "condition_cell": item.condition_cell or "",
            "condition_value": str(item.condition_value) if item.condition_value is not None else "",
        }
    with audit_path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows[key] for key in sorted(rows))
    draft.audit_rows=rows

    conversion_json = target.with_suffix(".conversion.json")
    if conversion_json.is_file():
        try:
            metadata = json.loads(conversion_json.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            metadata = {}
        metadata["last_project_edit"] = _edit_record(draft, target)
        _atomic_json(conversion_json, metadata)
    else:
        _atomic_json(target.with_suffix(".edit.json"), _edit_record(draft, target))


def _edit_record(draft: ExistingProjectDraft, target: Path) -> dict[str, Any]:
    return {
        "operation": "project_config_edit",
        "source_config": str(draft.source_path),
        "output_config": str(target),
        "target_version": draft.target_version,
        "case_id": draft.case_id,
        "edited_at_utc": datetime.now(timezone.utc).isoformat(),
        "changed_fields": [
            {"sheet": sheet, "parameter": parameter, "value": draft.get(sheet, parameter).value if parameter != '__rows__' else 'channel rows replaced'}
            for sheet, parameter in sorted(draft.changed)
        ],
    }


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    temp = path.with_name(path.name + ".tmp")
    temp.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    os.replace(temp, path)


__all__ = [
    "ExistingProjectDraft", "canonical_destination", "load_existing_project",
    "save_existing_project", "validate_existing_project",
]
