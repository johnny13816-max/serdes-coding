"""Reader for the corrected IEEE 802.3dj D3p2 200G CAKR 178 workbook.

The format identity describes the workbook schema, not one case's values.
Case 1 through Case 4 therefore use the same parser.
"""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from .schema import (
    AdhocConfigSnapshot,
    AdhocFormatDescriptor,
    AdhocSetting,
    ProbeResult,
    ProbeStatus,
)


FORMAT_ID = "ieee8023dj_com_adhoc/d3p2/200g_cakr_178/v1"


class IEEE8023djD3p2CAKR178V1Parser:
    """Parse one fixed family of official-style COM Ad Hoc workbooks."""

    descriptor = AdhocFormatDescriptor(
        format_id=FORMAT_ID,
        display_name="IEEE 802.3dj D3p2 200G CAKR 178 workbook",
        source_family="IEEE 802.3dj COM Ad Hoc",
        source_revision="D3p2",
        profile="200G CAKR 178",
        parser_revision=1,
        required_sheets=(
            "License Notice",
            "COM_Settings",
            "selections_and_tables",
            "keywords_20-Feb-2024",
        ),
        signature_keys=(
            "f_b",
            "Delta_f",
            "C_d",
            "L_s",
            "C_b",
            "R_0",
            "PKG_NAME",
            "z_p select",
            "L",
            "M",
            "c(-1)",
            "c(-2)",
            "c(1)",
            "f_r",
            "g_DC",
            "g_DC_HP",
            "f_z",
            "f_HP_PZ",
            "N_qb",
            "P_qc",
            "ffe_pre_tap_len",
            "ffe_post_tap_len",
            "N_bg",
            "N_bf",
            "N_f",
            "Local Search",
            "Non-zero Local Search Method",
            "Clip Method",
            "Include PCB",
        ),
        signature_blocks=(
            "PKG_LowR_CLASSA",
            "PKG_HiR_CLASSB",
            "PKG_Module",
            "PKG_Null",
        ),
    )

    _TOP_BLOCK_START_COLUMNS = (1, 6, 10, 15)
    _TOP_LAST_ROW = 32

    def probe(self, workbook_path: str | Path) -> ProbeResult:
        path = Path(workbook_path)
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets = set(workbook.sheetnames)
            matched: list[str] = []
            missing: list[str] = []

            for name in self.descriptor.required_sheets:
                signature = f"sheet:{name}"
                (matched if name in sheets else missing).append(signature)

            keys: set[str] = set()
            blocks: set[str] = set()
            if "COM_Settings" in sheets:
                keys, blocks = self._scan_signatures(workbook["COM_Settings"])

            for name in self.descriptor.signature_keys:
                signature = f"key:{name}"
                (matched if name in keys else missing).append(signature)
            for name in self.descriptor.signature_blocks:
                signature = f"block:{name}"
                (matched if name in blocks else missing).append(signature)

            ratio = len(matched) / max(1, len(matched) + len(missing))
            if not missing:
                status = ProbeStatus.EXACT
            elif ratio >= 0.8 and "sheet:COM_Settings" in matched:
                status = ProbeStatus.COMPATIBLE
            else:
                status = ProbeStatus.UNKNOWN
            return ProbeResult(
                status=status,
                descriptor=self.descriptor,
                matched_signatures=tuple(matched),
                missing_signatures=tuple(missing),
                candidates=(self.descriptor.format_id,),
            )
        finally:
            workbook.close()

    def parse(self, workbook_path: str | Path) -> AdhocConfigSnapshot:
        path = Path(workbook_path).expanduser().resolve()
        probe = self.probe(path)
        if probe.status not in {ProbeStatus.EXACT, ProbeStatus.COMPATIBLE}:
            missing = ", ".join(probe.missing_signatures)
            raise ValueError(
                f"Workbook is not compatible with {self.descriptor.format_id}; "
                f"missing signatures: {missing or 'unknown'}"
            )

        workbook = load_workbook(path, read_only=True, data_only=False)
        values_workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["COM_Settings"]
            values_sheet = values_workbook["COM_Settings"]
            settings = self._read_top_settings(sheet, values_sheet)
            settings.extend(self._read_named_blocks(sheet, values_sheet))
            return AdhocConfigSnapshot(
                descriptor=self.descriptor,
                source_path=path,
                source_sha256=self._sha256(path),
                settings=tuple(settings),
                sheet_names=tuple(workbook.sheetnames),
            )
        finally:
            workbook.close()
            values_workbook.close()

    @classmethod
    def _scan_signatures(cls, sheet: Any) -> tuple[set[str], set[str]]:
        keys: set[str] = set()
        for row in range(1, cls._TOP_LAST_ROW + 1):
            for start_col in cls._TOP_BLOCK_START_COLUMNS:
                value = sheet.cell(row, start_col + 1).value
                key = sheet.cell(row, start_col).value
                if key is not None and value is not None:
                    keys.add(str(key).strip())

        blocks: set[str] = set()
        for row in range(cls._TOP_LAST_ROW + 1, sheet.max_row + 1):
            if str(sheet.cell(row, 1).value).strip() == ".START":
                name = sheet.cell(row, 2).value
                if name is not None:
                    blocks.add(str(name).strip())
        return keys, blocks

    @classmethod
    def _read_top_settings(cls, sheet: Any, values_sheet: Any) -> list[AdhocSetting]:
        settings: list[AdhocSetting] = []
        for start_col in cls._TOP_BLOCK_START_COLUMNS:
            section = ""
            for row in range(1, cls._TOP_LAST_ROW + 1):
                key_cell = sheet.cell(row, start_col)
                value_cell = sheet.cell(row, start_col + 1)
                unit_cell = sheet.cell(row, start_col + 2)
                info_cell = sheet.cell(row, start_col + 3)
                key = cls._clean_key(key_cell.value)
                if key is None:
                    continue
                if value_cell.value is None:
                    section = key
                    continue
                settings.append(
                    AdhocSetting(
                        key=key,
                        raw_value=value_cell.value,
                        evaluated_value=values_sheet[value_cell.coordinate].value,
                        unit=unit_cell.value,
                        information=info_cell.value,
                        section=section,
                        block=None,
                        sheet=sheet.title,
                        key_cell=key_cell.coordinate,
                        value_cell=value_cell.coordinate,
                    )
                )
        return settings

    @classmethod
    def _read_named_blocks(cls, sheet: Any, values_sheet: Any) -> list[AdhocSetting]:
        settings: list[AdhocSetting] = []
        block: Optional[str] = None
        for row in range(cls._TOP_LAST_ROW + 1, sheet.max_row + 1):
            key_cell = sheet.cell(row, 1)
            value_cell = sheet.cell(row, 2)
            unit_cell = sheet.cell(row, 3)
            info_cell = sheet.cell(row, 4)
            key = cls._clean_key(key_cell.value)
            if key == ".START":
                block = None if value_cell.value is None else str(value_cell.value).strip()
                continue
            if key == ".END":
                block = None
                continue
            if block is None or key is None or key == "Parameter" or value_cell.value is None:
                continue
            settings.append(
                AdhocSetting(
                    key=key,
                    raw_value=value_cell.value,
                    evaluated_value=values_sheet[value_cell.coordinate].value,
                    unit=unit_cell.value,
                    information=info_cell.value,
                    section="package definition",
                    block=block,
                    sheet=sheet.title,
                    key_cell=key_cell.coordinate,
                    value_cell=value_cell.coordinate,
                )
            )
        return settings

    @staticmethod
    def _clean_key(value: Any) -> Optional[str]:
        if value is None:
            return None
        key = str(value).strip()
        return None if not key else key

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = sha256()
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()


__all__ = ["FORMAT_ID", "IEEE8023djD3p2CAKR178V1Parser"]
